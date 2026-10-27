import io
from unittest import mock

from django.test import SimpleTestCase

from common.documents import (
    DocumentCorruptError,
    DocumentEmptyError,
    DocumentTimeoutError,
    DocumentTooLargeError,
    DocumentUnsupportedError,
    get_parser,
)
from common.documents.constants import (
    CONTENT_TYPE_DOCX,
    CONTENT_TYPE_MARKDOWN,
    CONTENT_TYPE_PDF,
    CONTENT_TYPE_TEXT,
    MAX_FILE_BYTES,
)
from common.documents.docx_parser import DocxParser
from common.documents.fake import FakeDocumentParser
from common.documents.pdf_parser import PdfParser
from common.documents.text_parser import TextParser


class TextParserTests(SimpleTestCase):
    def test_text_and_markdown_pass_through(self):
        for content_type in (CONTENT_TYPE_TEXT, CONTENT_TYPE_MARKDOWN):
            parsed = TextParser().parse(b"\xef\xbb\xbf# Title\n\nBody\n", content_type)
            self.assertEqual(parsed.markdown, "# Title\n\nBody")
            self.assertEqual([page.number for page in parsed.pages], [None])

    def test_bad_bytes_and_nul_are_cleaned(self):
        parsed = TextParser().parse(b"caf\xe9\x00 ok", CONTENT_TYPE_TEXT)
        self.assertEqual(parsed.markdown, "caf� ok")

    def test_empty_unsupported_and_too_large(self):
        with self.assertRaises(DocumentEmptyError):
            TextParser().parse(b"  \n ", CONTENT_TYPE_TEXT)
        with self.assertRaises(DocumentUnsupportedError):
            TextParser().parse(b"x", CONTENT_TYPE_PDF)
        with (
            mock.patch("common.documents.text_parser.MAX_MARKDOWN_CHARS", 3),
            self.assertRaises(DocumentTooLargeError),
        ):
            TextParser().parse(b"abcd", CONTENT_TYPE_TEXT)


class FactoryAndFakeTests(SimpleTestCase):
    def test_routing(self):
        self.assertIsInstance(get_parser(CONTENT_TYPE_TEXT), TextParser)
        self.assertIsInstance(get_parser(CONTENT_TYPE_MARKDOWN), TextParser)
        self.assertIsInstance(get_parser(CONTENT_TYPE_PDF), PdfParser)
        self.assertIsInstance(get_parser(CONTENT_TYPE_DOCX), DocxParser)
        with self.assertRaises(DocumentUnsupportedError):
            get_parser("application/msword")

    def test_fake_parser_records_calls_and_fails_once(self):
        fake = FakeDocumentParser("# Hi")
        self.assertEqual(fake.parse(b"abc", CONTENT_TYPE_PDF).markdown, "# Hi")
        fake.fail_with = DocumentCorruptError()
        with self.assertRaises(DocumentCorruptError):
            fake.parse(b"abc", CONTENT_TYPE_PDF)
        fake.parse(b"abc", CONTENT_TYPE_PDF)
        self.assertEqual(len(fake.calls), 3)


def _docx_bytes(build) -> bytes:
    import docx

    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _heading_and_table(document):
    document.add_heading("Refund policy", level=1)
    document.add_paragraph("Refunds are issued within 14 days.")
    document.add_paragraph("Email us first", style="List Bullet")
    document.add_heading("Plans", level=2)
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Plan"
    table.rows[0].cells[1].text = "Price"
    table.rows[1].cells[0].text = "Pro"
    table.rows[1].cells[1].text = "$20 | month"


class DocxParserTests(SimpleTestCase):
    def test_headings_lists_and_tables_survive_in_order(self):
        markdown = (
            DocxParser()
            .parse(_docx_bytes(_heading_and_table), CONTENT_TYPE_DOCX)
            .markdown
        )
        self.assertLess(markdown.index("# Refund policy"), markdown.index("## Plans"))
        self.assertIn("- Email us first", markdown)
        self.assertIn("| Plan | Price |", markdown)
        self.assertIn("| Pro | $20 \\| month |", markdown)

    def test_docx_is_one_unnumbered_page(self):
        parsed = DocxParser().parse(_docx_bytes(_heading_and_table), CONTENT_TYPE_DOCX)
        self.assertEqual([page.number for page in parsed.pages], [None])

    def test_errors(self):
        with self.assertRaises(DocumentCorruptError):
            DocxParser().parse(b"PK\x03\x04 not really", CONTENT_TYPE_DOCX)
        with self.assertRaises(DocumentEmptyError):
            DocxParser().parse(_docx_bytes(lambda d: None), CONTENT_TYPE_DOCX)
        with self.assertRaises(DocumentUnsupportedError):
            DocxParser().parse(b"x", CONTENT_TYPE_PDF)
        with self.assertRaises(DocumentTooLargeError):
            DocxParser().parse(b"x" * (MAX_FILE_BYTES + 1), CONTENT_TYPE_DOCX)


def _pdf_bytes(pages: list[list[str]]) -> bytes:
    """A minimal text PDF: one list of lines per page (built by hand, no PDF library)."""
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        None,
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    kids = []
    for lines in pages:
        stream = (
            "BT /F1 12 Tf 72 720 Td 16 TL "
            + " ".join(f"({line}) Tj T*" for line in lines)
            + " ET"
        )
        page_id = len(objects) + 1
        kids.append(f"{page_id} 0 R")
        objects.append(
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {page_id + 1} 0 R >>"
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def _pdf_sized(lines: list[tuple[int, str]]) -> bytes:
    """A one-page text PDF where every line has its own font size: (size, text) pairs."""
    stream = "BT " + " ".join(
        f"/F1 {size} Tf 1 0 0 1 72 {740 - 22 * index} Tm ({text}) Tj"
        for index, (size, text) in enumerate(lines)
    )
    stream += " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [4 0 R] /Count 1 >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        (
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            "/Resources << /Font << /F1 3 0 R >> >> /Contents 5 0 R >>"
        ),
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


class PdfParserTests(SimpleTestCase):
    def test_larger_short_lines_become_headings_by_size(self):
        data = _pdf_sized(
            [
                (24, "Help Center"),
                (16, "Refunds"),
                (10, "You can request a refund within 14 days."),
                (10, "The money goes back to the card."),
                (16, "Shipping"),
                (10, "Orders ship in two days."),
            ]
        )
        markdown = PdfParser().parse(data, CONTENT_TYPE_PDF).pages[0].markdown
        self.assertIn("# Help Center", markdown)
        self.assertIn("\n## Refunds\n", markdown)
        self.assertIn("\n## Shipping\n", markdown)
        self.assertIn(
            "You can request a refund within 14 days.\nThe money goes back to the card.",
            markdown,
        )
        self.assertNotIn("# You can", markdown)

    def test_a_page_in_one_size_has_no_headings(self):
        data = _pdf_sized([(11, "Refunds"), (11, "Orders ship in two days.")])
        markdown = PdfParser().parse(data, CONTENT_TYPE_PDF).pages[0].markdown
        self.assertNotIn("#", markdown)

    def test_a_large_line_that_reads_like_a_sentence_is_not_a_heading(self):
        data = _pdf_sized(
            [(18, "This large line ends like a sentence."), (10, "Body.")]
        )
        markdown = PdfParser().parse(data, CONTENT_TYPE_PDF).pages[0].markdown
        self.assertNotIn("#", markdown)

    def test_pages_keep_their_numbers_and_empty_pages_are_dropped(self):
        data = _pdf_bytes([["Refunds take 14 days"], [], ["Contact support"]])
        parsed = PdfParser().parse(data, CONTENT_TYPE_PDF)
        self.assertEqual([page.number for page in parsed.pages], [1, 3])
        self.assertIn("Refunds take 14 days", parsed.pages[0].markdown)
        self.assertIn("Contact support", parsed.pages[1].markdown)

    def test_errors(self):
        with self.assertRaises(DocumentCorruptError):
            PdfParser().parse(b"%PDF-1.4 not really", CONTENT_TYPE_PDF)
        with self.assertRaises(DocumentEmptyError):
            PdfParser().parse(_pdf_bytes([[]]), CONTENT_TYPE_PDF)
        with self.assertRaises(DocumentUnsupportedError):
            PdfParser().parse(b"x", CONTENT_TYPE_DOCX)
        with self.assertRaises(DocumentTooLargeError):
            PdfParser().parse(b"x" * (MAX_FILE_BYTES + 1), CONTENT_TYPE_PDF)

    def test_page_cap_and_time_limit(self):
        data = _pdf_bytes([["a"], ["b"]])
        with (
            mock.patch("common.documents.pdf_parser.MAX_PAGES", 1),
            self.assertRaises(DocumentTooLargeError),
        ):
            PdfParser().parse(data, CONTENT_TYPE_PDF)
        with (
            mock.patch("common.documents.pdf_parser.PARSE_TIMEOUT_SECONDS", -1),
            self.assertRaises(DocumentTimeoutError),
        ):
            PdfParser().parse(data, CONTENT_TYPE_PDF)
